#!/usr/bin/env python3
"""RAG pipeline for an individual literature-gap matrix on low-cost visible spectroscopy.

This is a coursework prototype. It retrieves evidence per article, asks an LLM to
extract structured fields, and produces reviewable CSV/XLSX/JSON/Markdown outputs.
It does not measure spectra, validate a spectrometer, or establish novelty by itself.
"""
from __future__ import annotations
import argparse, json, os, re, sys, time
from pathlib import Path
from typing import Any

import fitz
import faiss
import numpy as np
import pandas as pd
from openai import OpenAI
from sentence_transformers import SentenceTransformer

EMBED_MODEL = "sentence-transformers/paraphrase-multilingual-mpnet-base-v2"
META_FIELDS = ["category", "title", "authors", "year", "venue", "doi", "url", "publication_status"]
EXTRACT_FIELDS = ["research_question", "domain_modality", "method", "rag_components", "evaluation", "key_findings", "limitations_reported", "relevance_to_spektra", "gap_evidence"]


def clean_text(s: str) -> str:
    s = re.sub(r"\s+", " ", s or "")
    return s.strip()


def chunk_text(text: str, size: int = 800, overlap: int = 120) -> list[str]:
    if size <= 0 or overlap < 0 or overlap >= size:
        raise ValueError("Require size > overlap >= 0")
    words = text.split()
    chunks=[]; start=0
    while start < len(words):
        end=min(len(words), start+size)
        chunks.append(" ".join(words[start:end]))
        if end == len(words): break
        start=end-overlap
    return chunks


def load_article_chunks(row: pd.Series, idx: int, size: int, overlap: int) -> list[dict[str, Any]]:
    result=[]
    pdf_path=str(row.get("pdf_path", "")).strip()
    if pdf_path:
        path=Path(pdf_path).expanduser()
        if not path.exists():
            print(f"WARNING: PDF not found for row {idx+1}: {path}; falling back to abstract/notes",file=sys.stderr)
        else:
            try:
                with fitz.open(path) as pdf:
                    for page_no,page in enumerate(pdf,1):
                        text=clean_text(page.get_text("text"))
                        for ci,chunk in enumerate(chunk_text(text,size,overlap)):
                            result.append({"article_idx":idx,"location":f"p.{page_no}","chunk":ci,"text":chunk})
            except Exception as e:
                print(f"WARNING: Could not read PDF {path}: {e}; falling back to abstract/notes",file=sys.stderr)
    if not result:
        text="\n".join(str(row.get(k,"")) for k in ["abstract","notes"] if str(row.get(k,"" )).strip())
        for ci,chunk in enumerate(chunk_text(clean_text(text),size,overlap)):
            result.append({"article_idx":idx,"location":"abstract/notes (no page-level full-text evidence)","chunk":ci,"text":chunk})
    return result


def parse_json(text: str) -> dict[str, Any]:
    text=re.sub(r"^```(?:json)?\s*|\s*```$", "", (text or "").strip(), flags=re.I)
    try:
        obj=json.loads(text)
    except json.JSONDecodeError:
        match=re.search(r"\{.*\}",text,re.S)
        if not match: raise
        obj=json.loads(match.group(0))
    if not isinstance(obj,dict): raise ValueError("LLM response was not a JSON object")
    return obj


def groq_client() -> OpenAI:
    key=os.getenv("GROQ_API_KEY","").strip()
    if not key: raise RuntimeError("GROQ_API_KEY is not set. Set it in the environment; do not hard-code or commit it.")
    return OpenAI(api_key=key,base_url=os.getenv("GROQ_BASE_URL","https://api.groq.com/openai/v1"))


def call_llm(client: OpenAI, model: str, system: str, prompt: str, temperature: float=0.1, retries: int=3) -> str:
    last=None
    for attempt in range(retries):
        try:
            res=client.chat.completions.create(model=model,temperature=temperature,messages=[{"role":"system","content":system},{"role":"user","content":prompt}])
            return res.choices[0].message.content or "{}"
        except Exception as exc:
            last=exc
            if attempt+1 < retries: time.sleep(2**attempt)
    raise RuntimeError(f"Groq request failed after {retries} attempts: {last}")


def main() -> None:
    ap=argparse.ArgumentParser(description="Build a literature gap matrix for RAG and a low-cost visible spectrometer.")
    ap.add_argument("--input",default="articles_individual_template.csv",help="CSV metadata template")
    ap.add_argument("--outdir",default="output")
    ap.add_argument("--model",default=os.getenv("GROQ_MODEL","llama-3.3-70b-versatile"),help="Groq model ID available to your account")
    ap.add_argument("--top-k",type=int,default=5)
    ap.add_argument("--chunk-size",type=int,default=800)
    ap.add_argument("--overlap",type=int,default=120)
    ap.add_argument("--skip-synthesis",action="store_true",help="Extract articles and matrix but skip cross-paper synthesis call")
    args=ap.parse_args()
    if args.top_k < 1: ap.error("--top-k must be >= 1")
    data=pd.read_csv(args.input,dtype=str,keep_default_na=False)
    required={"title","doi","abstract","notes","pdf_path"}
    missing=required-set(data.columns)
    if missing: ap.error(f"CSV missing columns: {', '.join(sorted(missing))}")
    if data.empty: ap.error("Input CSV has no article rows")
    out=Path(args.outdir);out.mkdir(parents=True,exist_ok=True)

    docs=[]
    for idx,row in data.iterrows(): docs.extend(load_article_chunks(row,idx,args.chunk_size,args.overlap))
    if not docs: ap.error("No text available; add abstracts/notes or readable local PDFs")
    embed=SentenceTransformer(EMBED_MODEL)
    vectors=embed.encode([d["text"] for d in docs],convert_to_numpy=True,normalize_embeddings=True,show_progress_bar=True).astype("float32")
    faiss_index=faiss.IndexFlatIP(vectors.shape[1]);faiss_index.add(vectors)
    client=groq_client()

    extracted_rows=[]; evidence=[]
    system=("You are a cautious scientific literature-review assistant. Use only supplied excerpts and metadata. "
            "Do not invent methods, results, limitations, page numbers, citations, or metrics. Distinguish reported facts from reviewer inference. "
            "RAG does not measure spectra. Preserve all values and units exactly. If not supported by excerpts, say 'Tidak dilaporkan pada teks yang diberikan'.")
    for idx,row in data.iterrows():
        meta={k:str(row.get(k,"")) for k in META_FIELDS}
        query=("Extract purpose, spectroscopy modality, RAG architecture, data, evaluation, findings and limitations. "
               "Assess relevance to an evidence-grounded assistant for a low-cost visible spectrometer (SPEKTRA), including data provenance, calibration, units, blank/reference, replicates and abstention only where discussed. "
               f"Article: {meta['title']} {row.get('abstract','')} {row.get('notes','')}")
        qv=embed.encode([query],convert_to_numpy=True,normalize_embeddings=True).astype("float32")[0]
        local=[(i,d) for i,d in enumerate(docs) if d["article_idx"]==idx]
        sims=[float(np.dot(qv,vectors[i])) for i,d in local]
        selected=[local[i][1] for i in np.argsort(sims)[::-1][:args.top_k]] if local else []
        context="\n\n".join(f"[SOURCE LOCATION: {d['location']} | CHUNK {d['chunk']}] {d['text']}" for d in selected)
        prompt=f"""Return exactly one valid JSON object with these keys: {', '.join(EXTRACT_FIELDS)}. For each key provide a concise evidence-based value. In gap_evidence, separate (a) what the article explicitly says and (b) cautious inference for this coursework. Include short supporting quotations and source locations when available. Do not state that something is absent from all literature; say only if absent from supplied excerpt.\n\nARTICLE METADATA:\n{json.dumps(meta,ensure_ascii=False)}\n\nEXCERPTS:\n{context}"""
        text=call_llm(client,args.model,system,prompt)
        extracted=parse_json(text)
        rec={**meta,"abstract_or_notes_only":not any(str(d["location"]).startswith("p.") for d in selected)}
        for field in EXTRACT_FIELDS:
            val=extracted.get(field,"Tidak dilaporkan pada teks yang diberikan")
            rec[field]=json.dumps(val,ensure_ascii=False) if not isinstance(val,str) else val
        extracted_rows.append(rec)
        evidence.append({"metadata":meta,"selected_evidence":[{"location":d["location"],"chunk":d["chunk"],"text":d["text"]} for d in selected],"extraction":extracted})
        print(f"Processed {idx+1}/{len(data)}: {meta['title']}")

    matrix=pd.DataFrame(extracted_rows)
    matrix.to_csv(out/"literature_gap_matrix.csv",index=False,encoding="utf-8-sig")
    with pd.ExcelWriter(out/"literature_gap_matrix.xlsx",engine="openpyxl") as writer:
        matrix.to_excel(writer,index=False,sheet_name="Gap Matrix")
        ws=writer.book["Gap Matrix"];ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
        from openpyxl.styles import Alignment,Font,PatternFill
        from openpyxl.utils import get_column_letter
        for c in ws[1]: c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="164A72");c.alignment=Alignment(wrap_text=True,vertical="top")
        for col in ws.columns:
            letter=get_column_letter(col[0].column);ws.column_dimensions[letter].width=min(max(max(len(str(x.value or "")) for x in col)*.65,14),38)
            for c in col[1:]:c.alignment=Alignment(wrap_text=True,vertical="top")
    (out/"article_evidence.json").write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding="utf-8")

    if not args.skip_synthesis:
        synthesis_input="\n\n".join(f"ARTICLE {i+1}: {json.dumps(e['extraction'],ensure_ascii=False)}\nMETADATA: {json.dumps(e['metadata'],ensure_ascii=False)}" for i,e in enumerate(evidence))
        synthesis_prompt=("Write a careful Indonesian synthesis for a coursework literature-gap matrix about an evidence-grounded RAG assistant for interpreting and documenting measurements from a low-cost visible spectrometer. "
          "Compare the core RAG/spectroscopy studies with supporting instrument/application references and cross-domain comparators. Do not treat non-RAG instrument papers as RAG studies. "
          "State that any gap is provisional within this reviewed corpus, not proof of absolute novelty. Distinguish reported evidence from proposed research. Propose a research question, measurable evaluation criteria, and limitations. Cite only author-year/DOI present in the supplied metadata.\n\nEVIDENCE:\n"+synthesis_input)
        analysis=call_llm(client,args.model,"You synthesize scientific literature cautiously and never overstate novelty.",synthesis_prompt,temperature=0.2)
    else:
        analysis="Sintesis tidak dibuat karena --skip-synthesis dipilih. Tinjau matriks dan susun analisis dengan membandingkan studi RAG/spektroskopi, sumber dasar instrumen, dan pembanding lintas domain."
    (out/"analysis.md").write_text("# Sintesis Literature Gap\n\n"+analysis+"\n\n## Pemeriksaan sebelum pengumpulan\n- Cocokkan metadata, DOI, metode, metrik, dan kutipan dengan artikel lengkap.\n- Bedakan hasil full text dengan hasil abstrak/catatan.\n- Gap ini hanya berlaku pada korpus yang ditelaah; perlu pencarian sistematis sebelum klaim novelty.\n- RAG membantu retrieval dan penjelasan; tidak menggantikan kalibrasi atau pengukuran spektrum.\n",encoding="utf-8")
    print(f"Done. Outputs saved to {out.resolve()}")

if __name__=="__main__": main()
